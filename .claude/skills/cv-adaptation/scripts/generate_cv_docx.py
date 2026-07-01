"""
Tool: generate_cv_docx.py
Responsabilidad: Genera un .docx del CV adaptado modificando el template ATS-optimizado.
Input:  <output_dir>/cv_{company}_{role}.json
Output: <output_dir>/cv_{company}_{role}.docx

Requiere que build_cv_template.py haya generado el template base previamente.
"""

import copy
import json
import re
import sys
import unicodedata
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


TEMPLATE_PATH    = Path(__file__).parents[4] / "assets" / "cv" / "cv_base_en.docx"
TEMPLATE_PATH_ES = Path(__file__).parents[4] / "assets" / "cv" / "cv_base_es.docx"
BULLET = "\u2022"


_ES_HEADERS = {
    "PROFESSIONAL SUMMARY":    "RESUMEN PROFESIONAL",
    "PROFESSIONAL EXPERIENCE": "EXPERIENCIA PROFESIONAL",
    "PROJECTS":                "PROYECTOS",
    "PERSONAL PROJECTS":       "PROYECTOS PERSONALES",
    "EDUCATION":               "FORMACIÓN",
    "SKILLS":                  "HABILIDADES TÉCNICAS",
    "TECHNICAL SKILLS":        "HABILIDADES TÉCNICAS",
    "LANGUAGES":               "IDIOMAS",
    "ADDITIONAL ACTIVITIES":   "ACTIVIDADES ADICIONALES",
}

# All known English→Spanish translations used as fallback when Spanish template is missing
_ES_HEADERS_FALLBACK = {
    **_ES_HEADERS,
    "EXPERIENCE":              "EXPERIENCIA PROFESIONAL",
    "CERTIFICATIONS":          "CERTIFICACIONES",
}


def generate_cv_docx(cv_adapted: dict, output_dir: str = "outputs") -> Path:
    """
    Genera el .docx del CV adaptado modificando el template base.

    Args:
        cv_adapted: dict con el CV adaptado (output de cv-adaptation)
        output_dir: directorio raíz de outputs (por defecto 'outputs')

    Returns:
        Path al archivo .docx generado
    """
    company  = cv_adapted.get("company", "UNKNOWN")
    role     = cv_adapted.get("role", "UNKNOWN")
    language = cv_adapted.get("language", "en")

    # Seleccionar template: español si existe y CV es en español, inglés otherwise
    using_es_template = (language == "es" and TEMPLATE_PATH_ES.exists())
    template_path = TEMPLATE_PATH_ES if using_es_template else TEMPLATE_PATH
    doc = Document(str(template_path))
    print(f"[generate_cv_docx] Usando template: {template_path.name}")

    # Traducir headers solo si caemos en el template inglés para un CV en español (fallback)
    if language == "es" and not using_es_template:
        _translate_section_headers(doc, _ES_HEADERS_FALLBACK)

    new_summary = cv_adapted.get("summary", "")
    if new_summary:
        _update_summary(doc, new_summary)

    experience_list = cv_adapted.get("experience", [])
    if experience_list:
        _update_experience(doc, experience_list)

    projects_list = cv_adapted.get("projects", [])
    if projects_list:
        _update_projects(doc, projects_list, language)

    education_list = cv_adapted.get("education", [])
    if education_list:
        _update_education(doc, education_list)

    if language == "es" and not using_es_template:
        _translate_section_headers_no_dict(doc)

    _remove_section(doc, _SECTIONS_TO_REMOVE)

    company_clean = re.sub(r"[^\w]", "_", company)
    role_clean    = re.sub(r"[^\w]", "_", role)[:30]

    output_path = Path(output_dir) / f"cv_{company_clean}_{role_clean}.docx"
    output_path.parent.mkdir(parents=True, exist_ok=True)

    doc.save(str(output_path))
    print(f"[generate_cv_docx] Guardado en: {output_path}")
    return output_path


# ── Helpers ────────────────────────────────────────────────────────────────────

def _normalize(text: str) -> str:
    """Lowercase + strip accents for fuzzy company name matching."""
    nfkd = unicodedata.normalize("NFD", text)
    return nfkd.encode("ascii", "ignore").decode("ascii").lower().strip()


_SUMMARY_HEADINGS = {"PROFESSIONAL SUMMARY", "RESUMEN PROFESIONAL"}

def _translate_section_headers(doc: Document, translations: dict[str, str]) -> None:
    """
    Traduce headers de sección del template en inglés al español.
    
    Args:
        doc: Document objeto de python-docx
        translations: dict mapping inglés → español para headers de sección
    """
    for p in doc.paragraphs:
        text_upper = p.text.strip().upper()
        if text_upper in translations:
            _replace_paragraph_text(p, translations[text_upper])
            print(f"[generate_cv_docx] Header traducido: '{text_upper}' → '{translations[text_upper]}'")

def _update_summary(doc: Document, new_summary: str) -> None:
    """
    Finds the paragraph immediately after the summary heading
    (English or Spanish) and replaces its text, preserving font formatting.
    """
    paragraphs = doc.paragraphs
    for i, p in enumerate(paragraphs):
        if p.text.strip().upper() in _SUMMARY_HEADINGS:
            for j in range(i + 1, min(i + 6, len(paragraphs))):
                candidate = paragraphs[j]
                if candidate.text.strip():
                    _replace_paragraph_text(candidate, new_summary)
                    return


def _update_experience(doc: Document, experience_list: list) -> None:
    """
    Locates company-line paragraphs (contain ' | ' AND a tab character)
    and replaces: (1) the job title/period in the company header line,
    (2) the subsequent bullet paragraphs (start with '•').

    Company-line pattern in template:
        "CompanyName | RoleTitle\\tDateRange"

    Matching key: first token of company name (before space or '('), normalized.
    Example: "Acme Corp (Parent)" → key "acme" → matches "Acme Corp | Data Scientist\\t..."
    """
    # Build lookup: normalized_first_token → {bullets, title, period, raw_name}
    company_data: dict[str, dict] = {}
    for exp in experience_list:
        raw_name = exp.get("company", "").strip()
        first_token = re.split(r"[\s(]", raw_name)[0]
        key = _normalize(first_token)
        bullets = exp.get("adapted_bullets", exp.get("bullets", []))
        if key:
            company_data[key] = {
                "bullets": bullets,
                "title": exp.get("title", ""),
                "period": exp.get("period", ""),
                "raw_name": raw_name,
            }

    # Keep backward-compat alias used below
    company_bullets: dict[str, list[str]] = {k: v["bullets"] for k, v in company_data.items()}

    if not company_bullets:
        return

    paragraphs = doc.paragraphs
    i = 0
    while i < len(paragraphs):
        p = paragraphs[i]
        p_text = p.text

        # Detect company line: has both ' | ' and a tab character
        if " | " in p_text and "\t" in p_text:
            company_segment = p_text.split(" | ")[0].strip()
            first_token = re.split(r"[\s(]", company_segment)[0]
            key = _normalize(first_token)

            if key in company_bullets:
                # Update job title/period in the company header line when provided
                entry = company_data.get(key, {})
                new_title  = entry.get("title", "")
                new_period = entry.get("period", "")
                raw_name   = entry.get("raw_name", company_segment)
                if new_title and new_period:
                    _replace_paragraph_text(p, f"{raw_name} | {new_title}\t{new_period}")
                elif new_title:
                    # Preserve existing period from template line
                    existing_period = p_text.split("\t", 1)[1] if "\t" in p_text else ""
                    _replace_paragraph_text(p, f"{raw_name} | {new_title}\t{existing_period}")

                # Collect indices of bullet paragraphs for this company block
                bullet_indices = []
                j = i + 1
                while j < len(paragraphs):
                    next_p = paragraphs[j]
                    next_text = next_p.text.strip()
                    # Stop at: next company line, section heading (all-caps), or end
                    if " | " in next_p.text and "\t" in next_p.text:
                        break
                    if (next_text and
                            next_text == next_text.upper() and
                            len(next_text) > 4 and
                            not next_text.startswith(BULLET)):
                        break
                    if next_text.startswith(BULLET):
                        bullet_indices.append(j)
                    j += 1

                _replace_bullet_block(doc, paragraphs, bullet_indices, company_bullets[key])
            else:
                # Not in adaptation JSON — normalize run structure for visual consistency
                # with adapted entries (collapses multi-run into run 0 which is bold=True).
                # Applies to experience entries omitted from the JSON and education entries.
                _replace_paragraph_text(p, p.text)

        i += 1


def _replace_paragraph_text(p, new_text: str, clear_bold: bool = False) -> None:
    """
    Replaces all text in a paragraph with new_text, preserving font of first run.
    Pass clear_bold=True for bullet paragraphs to prevent inheriting bold from template runs.
    """
    font_name = None
    font_size = None
    for run in p.runs:
        if run.text.strip():
            font_name = run.font.name
            font_size = run.font.size
            break

    for run in p.runs:
        run.text = ""

    if p.runs:
        r = p.runs[0]
        r.text = new_text
        if font_name:
            r.font.name = font_name
        if font_size:
            r.font.size = font_size
        if clear_bold:
            r.font.bold = False
    else:
        r = p.add_run(new_text)
        if font_name:
            r.font.name = font_name
        if font_size:
            r.font.size = font_size
        if clear_bold:
            r.font.bold = False


def _clean_bullet_text(text: str) -> str:
    """
    Removes space-hyphen-space used as a clause separator (AI-generation signal).
    Preserves compound hyphens in technical terms (out-of-time, K-means, end-to-end).
    """
    return re.sub(r" - ", ", ", text)


def _replace_bullet_block(
    doc: Document,
    paragraphs: list,
    bullet_indices: list,
    new_bullets: list,
) -> None:
    """
    Aligns existing bullet paragraphs with new_bullets.
    - Updates in-place for the min(len(existing), len(new)) bullets
    - Removes surplus existing bullets
    - Inserts new paragraphs (cloned from last bullet) when new_bullets > existing
    """
    n_existing = len(bullet_indices)
    n_new      = len(new_bullets)
    cleaned    = [_clean_bullet_text(b) for b in new_bullets]

    # Update in-place
    for k in range(min(n_existing, n_new)):
        p = paragraphs[bullet_indices[k]]
        _replace_paragraph_text(p, BULLET + "  " + cleaned[k], clear_bold=True)

    # Remove surplus existing bullets (iterate in reverse to preserve indices)
    for k in range(n_existing - 1, n_new - 1, -1):
        p_elem = paragraphs[bullet_indices[k]]._element
        p_elem.getparent().remove(p_elem)

    # Insert additional bullets when new_bullets > existing
    if n_new > n_existing and bullet_indices:
        anchor_elem = paragraphs[bullet_indices[-1]]._element
        for k in range(n_existing, n_new):
            new_elem = copy.deepcopy(anchor_elem)
            _set_first_run_text(new_elem, BULLET + "  " + cleaned[k], clear_bold=True)
            anchor_elem.addnext(new_elem)
            anchor_elem = new_elem


def _set_first_run_text(elem, text: str, clear_bold: bool = False) -> None:
    """Set text in the first w:r of an XML element, preserving run formatting.
    Clears w:t and w:tab from all runs to avoid leftover content from cloned paragraphs.
    Pass clear_bold=True for bullet elements to prevent inheriting bold from template runs."""
    runs = elem.findall(".//" + qn("w:r"))
    if not runs:
        return
    # Clear w:t and w:tab from ALL runs first
    for r_elem in runs:
        for child in list(r_elem):
            if child.tag in (qn("w:t"), qn("w:tab")):
                r_elem.remove(child)
    # Set new text only in the first run
    t = OxmlElement("w:t")
    t.text = text
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    runs[0].append(t)
    # Optionally strip bold from the first run's rPr
    if clear_bold:
        rPr = runs[0].find(qn("w:rPr"))
        if rPr is not None:
            for tag in (qn("w:b"), qn("w:bCs")):
                for elem_b in rPr.findall(tag):
                    rPr.remove(elem_b)


def _update_projects(doc: Document, projects_list: list, language: str = "en") -> None:
    """
    Inserts a PERSONAL PROJECTS section before EDUCATION if projects are present
    and the section doesn't already exist in the document.
    Each project gets a name/context header line and bullet(s).
    """
    if not projects_list:
        return

    edu_headings = {"EDUCATION", "FORMACIÓN", "EDUCACIÓN"}
    paragraphs = doc.paragraphs

    # Skip if PROJECTS/PERSONAL PROJECTS section already present
    for p in paragraphs:
        if p.text.strip().upper() in {"PROJECTS", "PROYECTOS", "PERSONAL PROJECTS", "PROYECTOS PERSONALES"}:
            return

    # Find EDUCATION heading element as insertion anchor
    edu_elem = None
    for p in paragraphs:
        if p.text.strip().upper() in edu_headings:
            edu_elem = p._element
            break

    if edu_elem is None:
        print("[generate_cv_docx] No se encontró sección EDUCATION para insertar PROJECTS")
        return

    # Find reference paragraph for section heading style (clone from EDUCATION heading)
    section_heading_ref = edu_elem

    # Find reference for entry header line (first " | " + "\t" paragraph)
    entry_line_ref = None
    for p in paragraphs:
        if " | " in p.text and "\t" in p.text:
            entry_line_ref = p._element
            break

    # Find reference for bullet paragraph style
    bullet_ref = None
    for p in paragraphs:
        if p.text.strip().startswith(BULLET):
            bullet_ref = p._element
            break

    if bullet_ref is None:
        print("[generate_cv_docx] No se encontró párrafo de bullet como referencia para PROJECTS")
        return

    section_title = "PROYECTOS PERSONALES" if language == "es" else "PERSONAL PROJECTS"

    # Collect new elements to insert (in document order)
    new_elements = []

    # PERSONAL PROJECTS section heading
    heading_elem = copy.deepcopy(section_heading_ref)
    _set_first_run_text(heading_elem, section_title)
    new_elements.append(heading_elem)

    # One entry per project
    for project in projects_list:
        name        = project.get("name", "")
        year        = str(project.get("year", ""))
        context     = project.get("context", "Personal Project")
        description = project.get("description", "")
        bullets     = project.get("bullets", [])

        # Entry header line: "Name | Context\tYear"
        ref = entry_line_ref if entry_line_ref is not None else bullet_ref
        entry_elem = copy.deepcopy(ref)
        _set_first_run_text(entry_elem, f"{name} | {context}\t{year}")
        new_elements.append(entry_elem)

        # Render bullets list (preferred) or fall back to single description
        if bullets:
            for b in bullets:
                b_elem = copy.deepcopy(bullet_ref)
                _set_first_run_text(b_elem, BULLET + "  " + _clean_bullet_text(b), clear_bold=True)
                new_elements.append(b_elem)
        elif description:
            desc_elem = copy.deepcopy(bullet_ref)
            _set_first_run_text(desc_elem, BULLET + "  " + _clean_bullet_text(description), clear_bold=True)
            new_elements.append(desc_elem)

    # Insert all elements before EDUCATION heading
    parent = edu_elem.getparent()
    edu_idx = list(parent).index(edu_elem)
    for offset, elem in enumerate(new_elements):
        parent.insert(edu_idx + offset, elem)

    print(f"[generate_cv_docx] Sección PERSONAL PROJECTS insertada con {len(projects_list)} proyecto(s)")


def _update_education(doc: Document, education_list: list) -> None:
    """
    Updates education entries in the doc from cv_adapted["education"].
    Matches by first token of institution name (normalized).
    Must run AFTER _update_experience() to avoid conflicts.
    """
    edu_data: dict[str, dict] = {}
    for edu in education_list:
        institution = edu.get("institution", "").strip()
        first_token = re.split(r"[\s(,]", institution)[0]
        key = _normalize(first_token)
        if key:
            edu_data[key] = edu

    if not edu_data:
        return

    paragraphs = doc.paragraphs
    i = 0
    while i < len(paragraphs):
        p = paragraphs[i]
        p_text = p.text

        if " | " in p_text and "\t" in p_text:
            # Try to match as education entry: institution is right of " | ", before tab
            parts = p_text.split(" | ", 1)
            if len(parts) == 2:
                right = parts[1].split("\t")[0].strip()
                first_token = re.split(r"[\s(,]", right)[0]
                key = _normalize(first_token)

                if key in edu_data:
                    edu = edu_data[key]
                    new_degree      = edu.get("degree", parts[0].strip())
                    new_institution = edu.get("institution", right)
                    new_period      = edu.get("period", edu.get("dates", ""))

                    _replace_paragraph_text(p, f"{new_degree} | {new_institution}\t{new_period}")
                    print(f"[generate_cv_docx] Educación actualizada: {new_degree[:40]}...")

                    # Next para is the location/description (italic line, not a company line)
                    new_location = edu.get("location", "")
                    if new_location and i + 1 < len(paragraphs):
                        loc_p = paragraphs[i + 1]
                        if not (" | " in loc_p.text and "\t" in loc_p.text):
                            _replace_paragraph_text(loc_p, new_location)

                    # Find note bullet within next 4 paragraphs
                    new_note = edu.get("note", "")
                    if new_note:
                        for j in range(i + 1, min(i + 5, len(paragraphs))):
                            candidate = paragraphs[j]
                            if candidate.text.strip().startswith(BULLET):
                                _replace_paragraph_text(candidate, BULLET + "  " + new_note)
                                break
                            if " | " in candidate.text and "\t" in candidate.text:
                                break

        i += 1


_SECTIONS_TO_REMOVE = {
    "ADDITIONAL ACTIVITIES",
    "ACTIVIDADES ADICIONALES",
}

def _remove_section(doc: Document, section_headers: set) -> None:
    """
    Removes a section (heading + all following paragraphs until next section or end)
    identified by its header text. Used to strip sections like ADDITIONAL ACTIVITIES.
    """
    paragraphs = doc.paragraphs
    start_idx = None
    for i, p in enumerate(paragraphs):
        if p.text.strip().upper() in section_headers:
            start_idx = i
            break

    if start_idx is None:
        return

    # Find end: next all-caps section heading or end of document
    end_idx = len(paragraphs)
    for i in range(start_idx + 1, len(paragraphs)):
        text = paragraphs[i].text.strip()
        if (text and text == text.upper() and len(text) > 3
                and not text.startswith(BULLET)):
            end_idx = i
            break

    header_text = paragraphs[start_idx].text.strip()
    body = doc.element.body
    for p in paragraphs[start_idx:end_idx]:
        body.remove(p._element)
    print(f"[generate_cv_docx] Sección eliminada: {header_text!r}")


def _translate_section_headers_no_dict(doc: Document) -> None:
    """
    Replaces all-caps English section headings with their Spanish equivalents.
    Called only when language == 'es'.
    """
    for p in doc.paragraphs:
        heading = p.text.strip().upper()
        if heading in _ES_HEADERS:
            _replace_paragraph_text(p, _ES_HEADERS[heading])


def load_json(path: str) -> dict:
    p = Path(path)
    if not p.exists():
        raise FileNotFoundError(f"Archivo no encontrado: {path}")
    with open(p, encoding="utf-8") as f:
        return json.load(f)


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Uso: python generate_cv_docx.py <cv_adapted.json> [output_dir]")
        print("Ejemplo: python generate_cv_docx.py outputs/ByRatings/cv_ByRatings_Data_Scientist.json")
        sys.exit(1)

    cv_json_path = sys.argv[1]
    output_dir   = sys.argv[2] if len(sys.argv) > 2 else "outputs"

    cv_adapted   = load_json(cv_json_path)
    result_path  = generate_cv_docx(cv_adapted, output_dir)
    print(f"[generate_cv_docx] Completado: {result_path}")
